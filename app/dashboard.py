import os
import secrets
import subprocess
import sys
from pathlib import Path

import httpx
from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, UploadFile, status
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, Response
from fastapi.security import HTTPBasic, HTTPBasicCredentials
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from sqlalchemy import delete, func, select

from app.config import settings
from app.models import Download, SessionLocal, Setting, User, init_db

app = FastAPI(title="Video Bot Dashboard")
security = HTTPBasic()
templates = Jinja2Templates(directory="app/templates")
ROOT = Path(__file__).resolve().parent.parent
try:
    app.mount("/static", StaticFiles(directory=str(ROOT / "app" / "static")), name="static")
except RuntimeError:
    pass


def _read_pid() -> int:
    try:
        return int((ROOT / "bot.pid").read_text(encoding="utf-8").strip())
    except (OSError, ValueError):
        return 0


def _bot_alive() -> int:
    pid = _read_pid()
    if not pid:
        return 0
    try:
        r = subprocess.run(["tasklist", "/FI", f"PID eq {pid}", "/NH"], capture_output=True, text=True, timeout=10)
    except Exception:
        return 0
    import re

    return pid if re.search(rf"\b{pid}\b", r.stdout) else 0


def _bot_start() -> str:
    if _bot_alive():
        return "already"
    log_out = open(ROOT / "bot_run.log", "a", encoding="utf-8")
    log_err = open(ROOT / "bot_run.err.log", "a", encoding="utf-8")
    p = subprocess.Popen([sys.executable, "-m", "app.bot"], cwd=str(ROOT), stdout=log_out, stderr=log_err)
    (ROOT / "bot.pid").write_text(str(p.pid), encoding="utf-8")
    return "started"


def _bot_stop() -> str:
    pid = _bot_alive()
    if not pid:
        killed = _kill_orphan_bots()
        try:
            (ROOT / "bot.pid").unlink()
        except OSError:
            pass
        return "stopped" if killed else "already_stopped"
    subprocess.run(["taskkill", "/PID", str(pid), "/F"], capture_output=True, timeout=15)
    _kill_orphan_bots()
    try:
        (ROOT / "bot.pid").unlink()
    except OSError:
        pass
    return "stopped"


def _kill_orphan_bots() -> bool:
    try:
        r = subprocess.run(
            ["powershell", "-NoProfile", "-Command",
             "Get-CimInstance Win32_Process -Filter \"Name = 'python.exe'\" | "
             "Where-Object { $_.CommandLine -like '*app.bot*' } | "
             "Select-Object -ExpandProperty ProcessId"],
            capture_output=True, text=True, timeout=30,
        )
    except Exception:
        return False
    killed = False
    for line in r.stdout.splitlines():
        line = line.strip()
        if line.isdigit():
            subprocess.run(["taskkill", "/PID", line, "/F"], capture_output=True, timeout=15)
            killed = True
    return killed


async def _resolve_token() -> str:
    cfg = await get_cfg()
    token = cfg.get("bot_token", "") or settings.bot_token
    if token == "put_telegram_bot_token_here":
        return ""
    return token


async def _auth_cfg() -> tuple:
    cfg = await get_cfg()
    user = cfg.get("dashboard_user", "") or settings.dashboard_user
    pw = cfg.get("dashboard_pass", "") or settings.dashboard_pass
    disabled = (cfg.get("auth_disabled", "") or os.getenv("AUTH_DISABLED", "0")) == "1"
    return user, pw, disabled


async def check_auth(request: Request):
    user, pw, disabled = await _auth_cfg()
    if disabled:
        return "open"
    try:
        creds = await security(request)
    except HTTPException:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="auth required",
                            headers={"WWW-Authenticate": "Basic"})
    if not (secrets.compare_digest(creds.username, user) and secrets.compare_digest(creds.password, pw)):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="auth required",
                            headers={"WWW-Authenticate": "Basic"})
    return creds.username


async def _save_login_fields(form, s) -> None:
    if "dashboard_user" in form and str(form["dashboard_user"]).strip():
        await _set_setting(s, "dashboard_user", str(form["dashboard_user"]).strip())
        write_env_key("DASHBOARD_USER", str(form["dashboard_user"]).strip())
    if "dashboard_pass_new" in form and str(form["dashboard_pass_new"]).strip():
        await _set_setting(s, "dashboard_pass", str(form["dashboard_pass_new"]).strip())
        write_env_key("DASHBOARD_PASS", str(form["dashboard_pass_new"]).strip())
    if "auth_disabled" in form:
        v = "1" if str(form["auth_disabled"]) == "1" else "0"
        await _set_setting(s, "auth_disabled", v)
        write_env_key("AUTH_DISABLED", v)


async def _set_setting(s, key: str, value: str) -> None:
    row = await s.get(Setting, key)
    if row is None:
        s.add(Setting(key=key, value=value))
    else:
        row.value = value


def mask_token(token: str) -> str:
    if not token or len(token) < 12:
        return ""
    return token[:6] + "..." + token[-4:]


def write_env_key(key: str, value: str) -> None:
    path = ".env"
    lines: list[str] = []
    found = False
    if os.path.exists(path):
        with open(path, "r", encoding="utf-8") as f:
            lines = f.readlines()
    for i, line in enumerate(lines):
        if line.strip().startswith(key + "="):
            lines[i] = f"{key}={value}\n"
            found = True
    if not found:
        lines.append(f"{key}={value}\n")
    with open(path, "w", encoding="utf-8") as f:
        f.writelines(lines)


def write_env_token(token: str) -> None:
    write_env_key("BOT_TOKEN", token)


async def telegram_get_me(token: str) -> dict:
    async with httpx.AsyncClient(timeout=15) as client:
        r = await client.get(f"https://api.telegram.org/bot{token}/getMe")
        return r.json()


async def get_cfg() -> dict[str, str]:
    async with SessionLocal() as s:
        rows = (await s.execute(select(Setting))).scalars().all()
        return {r.key: r.value for r in rows}


@app.on_event("startup")
async def startup():
    await init_db()


@app.get("/", response_class=HTMLResponse)
async def home(request: Request, user: str = Depends(check_auth)):
    try:
        page = max(1, int(request.query_params.get("page", "1")))
    except ValueError:
        page = 1
    per_page = 10
    async with SessionLocal() as s:
        users = (await s.execute(select(func.count(User.id)))).scalar() or 0
        downloads = (await s.execute(select(func.count(Download.id)))).scalar() or 0
        failed = (await s.execute(select(func.count(Download.id)).where(Download.status == "failed"))).scalar() or 0
        total_pages = max(1, (downloads + per_page - 1) // per_page)
        page = min(page, total_pages)
        last = (await s.execute(select(Download).order_by(Download.id.desc()).offset((page - 1) * per_page).limit(per_page))).scalars().all()
        clients = (await s.execute(select(User).order_by(User.id.desc()).limit(20))).scalars().all()
        counts: dict[int, int] = {}
        names: dict[int, str] = {}
        for c in clients:
            counts[c.id] = (await s.execute(select(func.count(Download.id)).where(Download.user_id == c.id))).scalar() or 0
            names[c.id] = ("@" + c.username) if c.username else str(c.telegram_id)
    cfg = await get_cfg()
    token = cfg.get("bot_token", "") or settings.bot_token
    bot_username = cfg.get("bot_username", "")
    bot_link = f"https://t.me/{bot_username}" if bot_username else ""
    cookies_file = cfg.get("cookies_file", "") or settings.cookies_file
    cookies_ok = bool(cookies_file and os.path.exists(cookies_file))
    return templates.TemplateResponse(
        "dashboard.html",
        {
            "request": request,
            "users": users,
            "downloads": downloads,
            "failed_count": failed,
            "last": last,
            "clients": clients,
            "counts": counts,
            "names": names,
            "page": page,
            "total_pages": total_pages,
            "bot_running": bool(_bot_alive()),
            "auth_user": (await _auth_cfg())[0],
            "auth_disabled": (await _auth_cfg())[2],
            "bcast": {"sent": request.query_params.get("b_sent", ""), "fail": request.query_params.get("b_fail", "")},
            "botmsg": request.query_params.get("botmsg", ""),
            "cfg": cfg,
            "token_masked": mask_token(token),
            "has_token": bool(token and token != "put_telegram_bot_token_here"),
            "bot_username": bot_username,
            "bot_link": bot_link,
            "cookies_file": cookies_file,
            "cookies_ok": cookies_ok,
        },
    )


@app.post("/settings")
async def save_settings(request: Request, user: str = Depends(check_auth)):
    form = await request.form()
    async with SessionLocal() as s:
        for k in ("welcome_en", "welcome_ar", "cookies_file"):
            if k in form:
                await _set_setting(s, k, str(form[k]).strip())
        await _save_login_fields(form, s)
        await s.commit()
    if "cookies_file" in form and str(form["cookies_file"]).strip():
        write_env_key("COOKIES_FILE", str(form["cookies_file"]).strip())
    return RedirectResponse("/", status_code=303)


@app.post("/token/save")
async def save_token(bot_token: str = Form(""), user: str = Depends(check_auth)):
    token = bot_token.strip()
    if len(token) < 20 or ":" not in token:
        raise HTTPException(status_code=400, detail="invalid token format")
    async with SessionLocal() as s:
        row = await s.get(Setting, "bot_token")
        if row is None:
            s.add(Setting(key="bot_token", value=token))
        else:
            row.value = token
        await s.commit()
    write_env_token(token)
    return RedirectResponse("/", status_code=303)


@app.post("/token/test", response_class=JSONResponse)
async def test_token(request: Request, user: str = Depends(check_auth)):
    data = await request.json() if request.headers.get("content-type", "").startswith("application/json") else {}
    token = (data.get("token") or "").strip()
    if not token:
        cfg = await get_cfg()
        token = cfg.get("bot_token", "") or settings.bot_token
    if not token or token == "put_telegram_bot_token_here":
        return JSONResponse({"ok": False, "error": "no token saved"})
    try:
        info = await telegram_get_me(token)
    except Exception as e:
        return JSONResponse({"ok": False, "error": str(e)[:200]})
    if not info.get("ok"):
        return JSONResponse({"ok": False, "error": str(info)[:300]})
    result = info["result"]
    username = result.get("username", "")
    link = f"https://t.me/{username}" if username else ""
    async with SessionLocal() as s:
        for k, v in {"bot_username": username, "bot_name": result.get("first_name", "")}.items():
            row = await s.get(Setting, k)
            if row is None:
                s.add(Setting(key=k, value=v))
            else:
                row.value = v
        await s.commit()
    return JSONResponse({"ok": True, "username": username, "name": result.get("first_name", ""), "link": link})


@app.post("/save-all")
async def save_all(request: Request, user: str = Depends(check_auth)):
    form = await request.form()
    token = str(form.get("bot_token", "")).strip()
    async with SessionLocal() as s:
        for k in ("welcome_en", "welcome_ar", "cookies_file"):
            if k in form:
                await _set_setting(s, k, str(form[k]).strip())
        if token:
            if len(token) < 20 or ":" not in token:
                raise HTTPException(status_code=400, detail="invalid token format")
            await _set_setting(s, "bot_token", token)
        await _save_login_fields(form, s)
        await s.commit()
    if token:
        write_env_key("BOT_TOKEN", token)
    if "cookies_file" in form and str(form["cookies_file"]).strip():
        write_env_key("COOKIES_FILE", str(form["cookies_file"]).strip())
    return RedirectResponse("/", status_code=303)


@app.get("/export")
async def export_all(user: str = Depends(check_auth)):
    from datetime import datetime
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    async with SessionLocal() as s:
        users = (await s.execute(select(User).order_by(User.id))).scalars().all()
        downloads = (await s.execute(select(Download).order_by(Download.id))).scalars().all()
        cfg = await get_cfg()
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"
    ws.append(["ID", "Telegram ID", "Username", "Lang", "Blocked", "Has login", "Downloads count", "Created at"])
    for u in users:
        ws.append([u.id, u.telegram_id, u.username, u.lang, "yes" if u.is_blocked else "no",
                   "yes" if u.cookies_path else "no", "",
                   u.created_at.strftime("%Y-%m-%d %H:%M") if u.created_at else ""])
    ws2 = wb.create_sheet("Downloads")
    ws2.append(["ID", "User ID", "URL", "Platform", "Type", "Status", "Error", "Error detail", "Created at"])
    per_user: dict[int, int] = {}
    for d in downloads:
        per_user[d.user_id] = per_user.get(d.user_id, 0) + 1
        ws2.append([d.id, d.user_id, d.url, d.platform, d.file_type, d.status, d.error_simple,
                    (d.error_detail or "")[:500],
                    d.created_at.strftime("%Y-%m-%d %H:%M") if d.created_at else ""])
    for row in ws.iter_rows(min_row=2):
        if row[0].value:
            row[6].value = per_user.get(row[0].value, 0)
    ws3 = wb.create_sheet("Settings")
    ws3.append(["Key", "Value"])
    for k, v in cfg.items():
        ws3.append([k, v if k != "bot_token" else (v[:6] + "..." if v else "")])
    for sheet in (ws, ws2, ws3):
        sheet.sheet_view.showGridLines = True
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 22
    buf = BytesIO()
    wb.save(buf)
    fname = "bot-data-" + datetime.utcnow().strftime("%Y%m%d-%H%M") + ".xlsx"
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@app.post("/cookies/upload")
async def cookies_upload(cookies: UploadFile = File(...), user: str = Depends(check_auth)):
    data = await cookies.read()
    if len(data) > 1_000_000:
        raise HTTPException(status_code=400, detail="file too large (max 1MB)")
    try:
        text = data.decode("utf-8", "ignore").lower()
    except Exception:
        text = ""
    if "youtube.com" not in text and "netscape" not in text:
        raise HTTPException(status_code=400, detail="not a valid cookies.txt export (no youtube.com found)")
    dest = ROOT / "cookies_youtube.txt"
    dest.write_bytes(data)
    async with SessionLocal() as s:
        row = await s.get(Setting, "cookies_file")
        if row is None:
            s.add(Setting(key="cookies_file", value=str(dest)))
        else:
            row.value = str(dest)
        await s.commit()
    write_env_key("COOKIES_FILE", str(dest))
    return RedirectResponse("/?botmsg=cookies_saved", status_code=303)


@app.post("/token/remove")
async def remove_token(user: str = Depends(check_auth)):
    async with SessionLocal() as s:
        for k in ("bot_token", "bot_username", "bot_name"):
            row = await s.get(Setting, k)
            if row is not None:
                row.value = ""
        await s.commit()
    write_env_key("BOT_TOKEN", "")
    _bot_stop()
    return RedirectResponse("/?botmsg=disconnected", status_code=303)


@app.post("/bot/start")
async def bot_start(user: str = Depends(check_auth)):
    return RedirectResponse(f"/?botmsg={_bot_start()}", status_code=303)


@app.post("/bot/stop")
async def bot_stop(user: str = Depends(check_auth)):
    return RedirectResponse(f"/?botmsg={_bot_stop()}", status_code=303)


@app.post("/broadcast")
async def broadcast(bmessage: str = Form(""), user: str = Depends(check_auth)):
    text = bmessage.strip()
    if not text:
        return RedirectResponse("/?b_sent=0&b_fail=0", status_code=303)
    token = await _resolve_token()
    if not token:
        return RedirectResponse("/?b_sent=0&b_fail=0", status_code=303)
    async with SessionLocal() as s:
        clients = (await s.execute(select(User).where(User.is_blocked.is_(False)))).scalars().all()
        tg_ids = [c.telegram_id for c in clients]
    sent, fail = 0, 0
    async with httpx.AsyncClient(timeout=15) as client:
        for tid in tg_ids:
            try:
                r = await client.post(f"https://api.telegram.org/bot{token}/sendMessage",
                                      json={"chat_id": tid, "text": text[:3000]})
                if r.json().get("ok"):
                    sent += 1
                else:
                    fail += 1
            except Exception:
                fail += 1
    return RedirectResponse(f"/?b_sent={sent}&b_fail={fail}", status_code=303)


@app.post("/users/{telegram_id}/block")
async def block_user(telegram_id: int, blocked: bool = Form(True), user: str = Depends(check_auth)):
    async with SessionLocal() as s:
        res = await s.execute(select(User).where(User.telegram_id == telegram_id))
        u = res.scalar_one_or_none()
        if not u:
            raise HTTPException(status_code=404, detail="user not found")
        u.is_blocked = bool(blocked)
        await s.commit()
    return RedirectResponse("/", status_code=303)


@app.get("/users", response_class=HTMLResponse)
async def users_list(request: Request, user: str = Depends(check_auth)):
    async with SessionLocal() as s:
        clients = (await s.execute(select(User).order_by(User.id.desc()).limit(200))).scalars().all()
        counts: dict[int, int] = {}
        for c in clients:
            n = (await s.execute(select(func.count(Download.id)).where(Download.user_id == c.id))).scalar() or 0
            counts[c.id] = n
    return templates.TemplateResponse("users.html", {"request": request, "clients": clients, "counts": counts})


@app.get("/users/{telegram_id}", response_class=HTMLResponse)
async def user_detail(request: Request, telegram_id: int, kind: str = "all", user: str = Depends(check_auth)):
    if kind not in ("all", "video", "audio"):
        kind = "all"
    async with SessionLocal() as s:
        res = await s.execute(select(User).where(User.telegram_id == telegram_id))
        client = res.scalar_one_or_none()
        if client is None:
            raise HTTPException(status_code=404, detail="user not found")
        q = select(Download).where(Download.user_id == client.id)
        if kind != "all":
            q = q.where(Download.file_type == kind)
        items = (await s.execute(q.order_by(Download.id.desc()).limit(200))).scalars().all()
    return templates.TemplateResponse(
        "user_detail.html", {"request": request, "client": client, "items": items, "kind": kind}
    )


@app.post("/history/clear")
async def history_clear(kind: str = Form("all"), user: str = Depends(check_auth)):
    if kind not in ("all", "video", "audio", "failed"):
        kind = "all"
    async with SessionLocal() as s:
        q = delete(Download)
        if kind == "failed":
            q = q.where(Download.status == "failed")
        elif kind in ("video", "audio"):
            q = q.where(Download.file_type == kind)
        await s.execute(q)
        await s.commit()
    return RedirectResponse("/", status_code=303)


@app.get("/api/errors")
async def api_errors(user: str = Depends(check_auth)):
    async with SessionLocal() as s:
        seen_row = await s.get(Setting, "errors_seen_id")
        seen = int(seen_row.value) if seen_row and (seen_row.value or "0").isdigit() else 0
        rows = (await s.execute(select(Download).where(Download.status == "failed").order_by(Download.id.desc()).limit(20))).scalars().all()
        out = []
        for r in rows:
            u = await s.get(User, r.user_id)
            out.append(
                {
                    "id": r.id,
                    "url": (r.url or "")[:80],
                    "platform": r.platform,
                    "kind": r.error_simple or "unknown",
                    "detail": (r.error_detail or "")[:400],
                    "at": r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "",
                    "client": ("@" + u.username) if u and u.username else (str(u.telegram_id) if u else "?"),
                }
            )
        unseen = sum(1 for r in rows if r.id > seen)
    return JSONResponse({"count": unseen, "errors": out})


@app.post("/api/errors/read")
async def api_errors_read(user: str = Depends(check_auth)):
    async with SessionLocal() as s:
        top = (await s.execute(select(func.max(Download.id)))).scalar() or 0
        row = await s.get(Setting, "errors_seen_id")
        if row is None:
            s.add(Setting(key="errors_seen_id", value=str(top)))
        else:
            row.value = str(top)
        await s.commit()
    return JSONResponse({"ok": True})
